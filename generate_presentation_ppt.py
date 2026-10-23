# generate_presentation_ppt.py
# This script creates a simple PowerPoint presentation (pptx) from the presentation text.
# Requires the `python-pptx` package.

presentation_text = """안녕하세요. 이번 발표에서는 대한민국 인프라 소외지역 분석 지도의 주요 기능과 데이터 분석 결과를 소개합니다.

- 전국 18대 인프라(에너지, 교통, 금융 등) 실시간 상황 시각화
- 위성 고해상도 이미지와 행정구역 검색 기능 제공
- 사용자 맞춤 레이어 토글 및 위험 영역 자동 표시"""

def create_pptx(output_path="presentation.pptx"):
    from pptx import Presentation
    from pptx.util import Inches, Pt
    prs = Presentation()
    # Title slide
    title_slide_layout = prs.slide_layouts[0]
    slide = prs.slides.add_slide(title_slide_layout)
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "프로젝트 발표"
    subtitle.text = "대한민국 인프라 소외지역 분석 지도"
    # Content slide
    bullet_slide_layout = prs.slide_layouts[1]
    slide2 = prs.slides.add_slide(bullet_slide_layout)
    title2 = slide2.shapes.title
    body = slide2.shapes.placeholders[1].text_frame
    title2.text = "발표 내용"
    # Split lines
    lines = presentation_text.strip().split('\n')
    for line in lines:
        if line.startswith('-'):
            p = body.add_paragraph()
            p.text = line[1:].strip()
            p.level = 0
        else:
            # treat as normal paragraph
            p = body.add_paragraph()
            p.text = line.strip()
            p.level = 0
    prs.save(output_path)
    print(f"Saved PPTX to {output_path}")

if __name__ == "__main__":
    create_pptx()
